from io import BytesIO

from docx import Document

from docx.enum.text import (
    WD_ALIGN_PARAGRAPH
)

from docx.enum.table import (
    WD_TABLE_ALIGNMENT
)

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from docx.shared import (
    Inches,
    Pt
)

from document_utils.sanitize import (
    sanitize_text
)


def _set_cell_text(
    cell,
    text: str,
    bold: bool = False
):

    cell.text = ""

    paragraph = cell.paragraphs[0]

    run = paragraph.add_run(
        text
    )

    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)


def _set_repeat_table_header(row):

    tr_pr = row._tr.get_or_add_trPr()

    tbl_header = OxmlElement(
        "w:tblHeader"
    )

    tbl_header.set(
        qn("w:val"),
        "true"
    )

    tr_pr.append(
        tbl_header
    )


def format_docx(
    text: str,
    doc_type: str,
    terms: str = ""
) -> bytes:

    doc = Document()

    section = doc.sections[0]

    section.top_margin = Inches(
        0.75
    )

    section.bottom_margin = Inches(
        0.75
    )

    section.left_margin = Inches(
        0.85
    )

    section.right_margin = Inches(
        0.85
    )

    normal_style = doc.styles[
        "Normal"
    ]

    normal_style.font.name = (
        "Times New Roman"
    )

    normal_style.font.size = Pt(
        11
    )

    # Brand
    p = doc.add_paragraph()

    p.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = p.add_run(
        "LEGALEASE"
    )

    run.bold = True
    run.font.name = (
        "Times New Roman"
    )

    run.font.size = Pt(
        18
    )

    # Document title
    p = doc.add_paragraph()

    p.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = p.add_run(
        sanitize_text(doc_type)
    )

    run.bold = True
    run.font.name = (
        "Times New Roman"
    )

    run.font.size = Pt(
        14
    )

    clean = sanitize_text(
        text
    )

    for block in clean.split(
        "\n\n"
    ):

        block = block.strip()

        if not block:
            continue

        lines = block.splitlines()

        # Heading
        if (
            len(lines) == 1
            and len(lines[0]) <= 100
            and (
                lines[0].isupper()
                or lines[0].endswith(":")
            )
        ):

            p = doc.add_paragraph()

            run = p.add_run(
                lines[0]
            )

            run.bold = True

            run.font.name = (
                "Times New Roman"
            )

            run.font.size = Pt(
                12
            )

            continue

        # Bullets
        if all(
            line.lstrip().startswith(
                ("-", "*")
            )
            for line in lines
        ):

            for line in lines:

                p = doc.add_paragraph(
                    style="List Bullet"
                )

                p.add_run(
                    line.lstrip()[1:].strip()
                )

            continue

        # Normal paragraph
        p = doc.add_paragraph()

        for index, line in enumerate(
            lines
        ):

            if index:
                p.add_run("\n")

            p.add_run(
                line
            )

    # Terms table
    if terms.strip():

        term_items = [
            item.strip()
            for item in terms.split(";")
            if item.strip()
        ]

        if term_items:

            doc.add_page_break()

            heading = doc.add_paragraph()

            run = heading.add_run(
                "KEY TERMS"
            )

            run.bold = True
            run.font.name = (
                "Times New Roman"
            )

            run.font.size = Pt(
                12
            )

            table = doc.add_table(
                rows=1,
                cols=2
            )

            table.alignment = (
                WD_TABLE_ALIGNMENT.CENTER
            )

            table.style = (
                "Table Grid"
            )

            _set_cell_text(
                table.rows[0].cells[0],
                "No.",
                True
            )

            _set_cell_text(
                table.rows[0].cells[1],
                "Term / Condition",
                True
            )

            _set_repeat_table_header(
                table.rows[0]
            )

            for number, item in enumerate(
                term_items,
                1
            ):

                row = table.add_row()

                _set_cell_text(
                    row.cells[0],
                    str(number)
                )

                _set_cell_text(
                    row.cells[1],
                    item
                )

    # Footer
    footer = (
        section.footer
        .paragraphs[0]
    )

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = footer.add_run(
        "LegalEase - AI-assisted draft. "
        "Review before legal use."
    )

    run.font.name = (
        "Times New Roman"
    )

    run.font.size = Pt(
        8
    )

    buffer = BytesIO()

    doc.save(
        buffer
    )

    return buffer.getvalue()