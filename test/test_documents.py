from io import BytesIO

from docx import Document

from document_utils.docx_generator import (
    format_docx
)

from document_utils.pdf_generator import (
    format_pdf
)

from document_utils.txt_generator import (
    format_txt
)


TEXT = """
SERVICE AGREEMENT

The parties agree to the following terms.

- Payment is due within 30 days.
- Confidentiality applies to all non-public information.
"""


def test_txt_export():

    result = format_txt(
        TEXT
    )

    assert isinstance(
        result,
        bytes
    )

    assert (
        b"SERVICE AGREEMENT"
        in result
    )


def test_docx_export():

    result = format_docx(
        TEXT,
        "Service Agreement",
        "Payment is due within 30 days;"
        "Confidentiality applies"
    )

    assert result[:2] == b"PK"

    document = Document(
        BytesIO(result)
    )

    joined = "\n".join(
        paragraph.text
        for paragraph
        in document.paragraphs
    )

    assert (
        "SERVICE AGREEMENT"
        in joined
    )

    assert len(
        document.tables
    ) == 1


def test_pdf_export():

    result = format_pdf(
        TEXT,
        "Service Agreement"
    )

    assert result.startswith(
        b"%PDF"
    )

    assert len(result) > 1000